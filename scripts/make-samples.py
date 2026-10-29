"""Generate entirely fictional DOCX fixtures; no original documents are used."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
ROOT=Path(__file__).resolve().parents[1]
def fresh(title):
 d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.75)
 normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(8)
 for name in ['Title','Heading 1','Heading 2']:
  d.styles[name].font.color.rgb=RGBColor(0,0,0)
 for style in d.styles:
  for border in style.element.xpath('.//w:pBdr'):
   border.getparent().remove(border)
 d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title=title;d.core_properties.subject='Fictional publishing demonstration';d.core_properties.comments=''
 return d

d=fresh('A morning on the ridge')
d.add_heading('A morning on the ridge',0)
d.add_paragraph('<div id=datestamp>FIELD JOURNAL / FICTIONAL SAMPLE / OCTOBER 2026</div>')
d.add_paragraph('This fictional field journal demonstrates how a Word document can become a web article. Ordinary paragraphs stay readable, while a few simple markers become reusable components.')
d.add_paragraph('The trail began quietly, beyond the last wooden gate. A pale line of sunlight touched the ridge, and the valley below was still folded into shade. We stopped for a moment before following the path uphill.')
d.add_paragraph('<img id=ridge />')
d.add_heading('Room to notice',1)
d.add_paragraph('The route was not especially long. Its best moments were small ones: a bend in the trail, a patch of warm stone, and the changing colors of the hills. The journal kept those observations in one place.')
d.add_paragraph('<div id=quote>')
d.add_paragraph('Some walks are measured in miles. Others are measured in the things you finally had time to notice.')
d.add_paragraph('</div>')
d.add_paragraph('<div id=info>')
d.add_paragraph('Publishing note: this box comes from an info marker typed directly into Word. The source document contains text, not a screenshot of a designed page.')
d.add_paragraph('</div>')
d.add_heading('A simple packing list',1)
for item in ['A notebook for the details worth keeping','A water bottle and a light layer','Enough time to take the slower route']:
 d.add_paragraph(item,style='List Bullet')
d.add_paragraph('<div id=horizontalLine></div>')
d.add_heading('Back at the desk',1)
d.add_paragraph('The same document can be edited, saved, and converted again. That is the small idea behind this demo: keep writing comfortable, then let components handle the presentation.')
d.save(ROOT/'public/samples/Field-Journal.docx')
f=fresh('Component reference')
f.add_heading('Component reference',0)
f.add_paragraph('Neutral test content for the DOCX publishing pipeline.')
for text in ['<div id=quote>','A sample quotation with ordinary words.','</div>','<div id=info>','A sample information panel.','</div>','<img id=ridge />','<div id=datestamp>FICTIONAL SAMPLE</div>','<div id=horizontalLine></div>','<div id=video></div>']:
 f.add_paragraph(text)
f.save(ROOT/'public/samples/Component-Reference.docx')
